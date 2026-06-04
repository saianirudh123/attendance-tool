"""
Build the attendance register workbook from parsed data.
Returns a summary dict for the dashboard.
"""
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import date, datetime, time, timedelta
import calendar
import re

# ── Styles ────────────────────────────────────────────────────────────────────
def _font(bold=False, color="000000", size=10):
    return Font(bold=bold, color=color, name="Calibri", size=size)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

HDR_FILL   = _fill("1F3864"); HDR_FONT  = _font(True, "FFFFFF", 10)
LBL_FILL   = _fill("EBF3FB"); LBL_FONT  = _font(True, "1F3864", 10)
PRESENT_F  = _fill("C6EFCE"); ABSENT_F  = _fill("FFC7CE")
WO_F       = _fill("FFEB9C"); HOL_F     = _fill("E2EFDA")
CLOSE_F    = _fill("D9EAD3"); LWP_F     = _fill("F4CCCC")
CENTER     = Alignment(horizontal="center", vertical="center")
LEFT       = Alignment(horizontal="left",   vertical="center")
THIN_B     = _border()

def _hdr(cell, text):
    cell.value = text; cell.fill = HDR_FILL; cell.font = HDR_FONT
    cell.alignment = CENTER; cell.border = THIN_B

def _lbl(cell, text):
    cell.value = text; cell.fill = LBL_FILL; cell.font = LBL_FONT
    cell.alignment = LEFT; cell.border = THIN_B

def _dat(cell, val=None, fmt=None, bold=False, fill=None):
    if val is not None: cell.value = val
    cell.font = _font(bold); cell.alignment = CENTER; cell.border = THIN_B
    if fmt:  cell.number_format = fmt
    if fill: cell.fill = fill

def _set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def _t(hhmm):
    """'HH:MM' -> Excel time fraction"""
    if not hhmm: return None
    h, m = map(int, hhmm.split(":"))
    return (h * 60 + m) / 1440.0

def _day_type(d: date, week_off_day: int, holidays: list) -> str:
    """week_off_day: 0=Mon … 6=Sun (per-employee, read from Excel)."""
    if d in holidays: return "Company Holiday"
    return "Week Off" if d.weekday() == week_off_day else "Working Day"

def _resolve_code(emp_name: str, name_to_code: dict) -> str:
    key = emp_name.lower()
    if key in name_to_code:
        return name_to_code[key]
    tokens = key.split()
    for pdf_name, code in name_to_code.items():
        matches = sum(1 for t in tokens if t in pdf_name.split())
        if matches >= 2:
            return code
        if tokens and tokens[-1] in pdf_name.split():
            return code
    return None


# ── Computation helpers (mirror Excel formulas in Python for JSON output) ─────

def _time_to_mins(t: str):
    """'HH:MM' → minutes from midnight, or None."""
    if not t:
        return None
    try:
        h, m = map(int, t.split(":"))
        return h * 60 + m
    except Exception:
        return None

def _punch_duration_mins(in_t, out_t) -> int:
    """Return duration minutes for a punch pair; out before in means next day."""
    in_m = _time_to_mins(in_t)
    out_m = _time_to_mins(out_t)
    if in_m is None or out_m is None:
        return 0
    return out_m - in_m if out_m >= in_m else (24 * 60 - in_m) + out_m

def _extra_punch_mins(extra_punches) -> int:
    return sum(_punch_duration_mins(p.get("in_time"), p.get("out_time"))
               for p in (extra_punches or []))

def _extra_punch_label(extra_punches) -> str:
    labels = []
    for p in (extra_punches or []):
        in_t = p.get("in_time") or ""
        out_t = p.get("out_time") or ""
        if in_t or out_t:
            labels.append(f"{in_t or '?'}-{out_t or '?'}")
    return "\n".join(labels)

def _duration_to_mins(v) -> int:
    if v in (None, ""):
        return 0
    if isinstance(v, timedelta):
        return int(round(v.total_seconds() / 60))
    if isinstance(v, (int, float)):
        return int(round(float(v) * 1440))
    if isinstance(v, time):
        return v.hour * 60 + v.minute
    if isinstance(v, str):
        parsed = _time_to_mins(v)
        return parsed or 0
    return 0

def _cell_time_to_str(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, time):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, timedelta):
        mins = _duration_to_mins(v)
        return f"{(mins // 60) % 24:02d}:{mins % 60:02d}"
    if isinstance(v, (int, float)):
        mins = _duration_to_mins(v)
        return f"{(mins // 60) % 24:02d}:{mins % 60:02d}"
    if isinstance(v, str):
        s = v.strip()
        if re.match(r"^\d{1,2}:\d{2}$", s):
            h, m = s.split(":")
            return f"{int(h):02d}:{int(m):02d}"
    return None

def _parse_extra_punches(v):
    if not v:
        return []
    punches = []
    for part in str(v).replace(",", "\n").splitlines():
        m = re.match(r"\s*(\d{1,2}:\d{2}|\?)\s*[-–]\s*(\d{1,2}:\d{2}|\?)\s*", part)
        if not m:
            continue
        in_t = None if m.group(1) == "?" else _cell_time_to_str(m.group(1))
        out_t = None if m.group(2) == "?" else _cell_time_to_str(m.group(2))
        punches.append({"in_time": in_t, "out_time": out_t})
    return punches

def _mins_to_str(mins) -> str:
    """Minutes → 'H:MM' string (supports negative)."""
    if mins is None:
        return None
    sign = "-" if mins < 0 else ""
    mins = abs(int(round(mins)))
    return f"{sign}{mins // 60}:{mins % 60:02d}"

def _get_thresholds(location: str) -> dict:
    """Return location-specific minute thresholds."""
    if "factory" in location.lower():
        return dict(sh_in=9*60+15, sh_out=17*60+45,
                    lh_in=8*60+30, lh_out=18*60+45,
                    lm_lo=10*60,   lm_hi=13*60+30,
                    hd_thr=13*60+30)
    return dict(sh_in=9*60+45, sh_out=18*60+15,
                lh_in=9*60+15, lh_out=19*60+15,
                lm_lo=10*60+30, lm_hi=14*60,
                hd_thr=14*60)

def _compute_day_row(emp_type: str, in_t, out_t, day_type: str,
                     location: str, extra_punches=None) -> dict:
    """Compute all daily metrics for one day."""
    th = _get_thresholds(location)
    in_m  = _time_to_mins(in_t)
    out_m = _time_to_mins(out_t)
    extra_mins = _extra_punch_mins(extra_punches)
    has_extra = extra_mins > 0 or any(p.get("in_time") or p.get("out_time")
                                      for p in (extra_punches or []))
    status = "Present" if (in_t or out_t or has_extra) else "Absent"

    if emp_type == "Labour":
        hw = (out_m - in_m) if (status == "Present" and in_m is not None and out_m is not None) else 0
        hw += extra_mins
        ot = (hw - 9 * 60) if status == "Present" else 0
        return dict(status=status, short_mins=0, long_mins=0,
                    net_mins=hw, overtime_mins=ot, late_mark="", half_day="",
                    extra_mins=extra_mins)

    # Full-time
    half_day = ""
    if day_type == "Working Day" and status == "Present":
        if in_m is not None and in_m >= th["hd_thr"]:
            half_day = "Half Day"
        elif out_m is not None and out_m <= th["hd_thr"]:
            half_day = "Half Day"

    short_mins = long_mins = net_mins = 0
    if not half_day:
        if day_type == "Working Day" and status == "Present" and in_m is not None and out_m is not None:
            short_mins = max(0, in_m - th["sh_in"]) + max(0, th["sh_out"] - out_m)
            long_mins  = max(0, th["lh_in"] - in_m) + max(0, out_m - th["lh_out"])
            net_mins   = long_mins - short_mins
        elif day_type in ("Company Holiday", "Week Off") and status == "Present" and in_m is not None and out_m is not None:
            net_mins = out_m - in_m
    net_mins += extra_mins

    late_mark = ""
    if day_type == "Working Day" and status == "Present" and in_m is not None:
        if th["lm_lo"] <= in_m < th["lm_hi"]:
            late_mark = "Late Mark"

    return dict(status=status, short_mins=short_mins, long_mins=long_mins,
                net_mins=net_mins, overtime_mins=0, late_mark=late_mark,
                half_day=half_day, extra_mins=extra_mins)


def _compute_emp_summary(emp: dict, daily_rows: list) -> dict:
    """Compute monthly + leave summary from daily rows (mirrors Excel formulas)."""
    wd_present = [d for d in daily_rows if d["day_type"] == "Working Day" and d["status"] == "Present"]
    wd_absent  = [d for d in daily_rows if d["day_type"] == "Working Day" and d["status"] == "Absent"]
    half_days  = [d for d in daily_rows if d["half_day"] == "Half Day"]
    total_net_mins = sum(d["net_mins"] for d in daily_rows)

    present_days  = len(wd_present)
    absent_days   = len(wd_absent)
    num_half_days = len(half_days)

    if emp["type"] == "Labour":
        all_present = [d for d in daily_rows if d["status"] == "Present"]
        total_ot_mins = sum(d["overtime_mins"] for d in daily_rows)
        working_days  = sum(1 for d in daily_rows if d["day_type"] == "Working Day")
        return dict(present_days=len(all_present), absent_days=absent_days,
                    half_days=0, total_net_hours=round(total_net_mins / 60, 4),
                    working_days=working_days,
                    total_overtime_hours=round(total_ot_mins / 60, 4),
                    ot_days=round(total_ot_mins / (6 * 60), 4),
                    el_opening=0, cl_opening=0, sl_opening=0, sl_applied=0)

    el_o = emp["el"] or 0
    cl_o = emp["cl"] or 0
    sl_o = emp["sl"] or 0
    sl_a = emp["sl_applied"] or 0

    net_hours  = total_net_mins / 60
    el_accrual = math.ceil(net_hours / 9) if net_hours < 0 else math.floor(net_hours / 9)

    total_absent  = absent_days + 0.5 * num_half_days
    cl_adj        = min(1, cl_o, max(0, total_absent))
    sl_adj        = min(sl_a, sl_o)
    adj_el_open   = max(0, el_o + el_accrual)
    el_adj        = min(adj_el_open, max(0, total_absent - cl_adj - sl_adj))
    lwp_neg_el    = max(0, -(el_o + el_accrual)) if (el_o + el_accrual) < 0 else 0
    lwp_excess    = max(0, total_absent - cl_adj - sl_adj - el_adj)
    total_lwp     = lwp_neg_el + lwp_excess

    return dict(present_days=present_days, absent_days=absent_days,
                half_days=num_half_days,
                total_net_hours=round(net_hours, 4),
                el_opening=el_o, cl_opening=cl_o, sl_opening=sl_o, sl_applied=sl_a,
                el_accrual=el_accrual,
                cl_adj=cl_adj, sl_adj=sl_adj, el_adj=el_adj,
                total_lwp=round(total_lwp, 4),
                cl_closing=max(0, cl_o - cl_adj),
                el_closing=max(0, adj_el_open - el_adj),
                sl_closing=max(0, sl_o - sl_adj))


def _build_emp_json(emp: dict, punch_data: dict, name_to_code: dict,
                    holidays: list, year: int, month: int) -> dict:
    """Return enriched employee data dict for frontend rendering."""
    loc      = emp["location"]
    name     = emp["name"]
    week_off = emp["week_off"]
    code     = _resolve_code(name, name_to_code)
    punch    = punch_data.get(code, {}) if code else {}
    days_in_month = calendar.monthrange(year, month)[1]

    daily_rows = []
    for day_num in range(1, days_in_month + 1):
        d        = date(year, month, day_num)
        day_type = _day_type(d, week_off, holidays)
        in_t, out_t = punch.get(day_num, (None, None))
        metrics  = _compute_day_row(emp["type"], in_t, out_t, day_type, loc, [])
        daily_rows.append(dict(day=day_num, date=d.strftime("%d-%b"),
                               dow=d.strftime("%a"), day_type=day_type,
                               in_time=in_t, out_time=out_t,
                               extra_punches=[], **metrics))

    summary = _compute_emp_summary(emp, daily_rows)
    punch_days = len([v for v in punch.values() if v[0] or v[1]])

    return dict(name=name, code=code or "?",
                emp_id=emp.get("emp_id", ""),
                designation=emp.get("designation", ""),
                department=emp.get("department", ""),
                area=emp.get("area", ""),
                type=emp["type"], location=loc,
                week_off=week_off, punch_days=punch_days, daily=daily_rows, **summary)


# ── Employee sheet ─────────────────────────────────────────────────────────────
def _build_employee_sheet(wb, emp: dict, punch_data: dict, name_to_code: dict,
                          holidays: list, year: int, month: int):
    loc        = emp["location"]
    is_factory = "factory" in loc.lower()
    week_off   = emp["week_off"]
    name       = emp["name"]
    ws         = wb.create_sheet(title=name[:28])
    ws.freeze_panes = "A4"

    code  = _resolve_code(name, name_to_code)
    punch = punch_data.get(code, {}) if code else {}
    days_in_month = calendar.monthrange(year, month)[1]

    emp_id  = emp.get("emp_id", "")
    desig   = emp.get("designation", "")
    dept    = emp.get("department", "")
    area_v  = emp.get("area", "")

    ws.merge_cells("A1:M1")
    t = ws["A1"]
    t.value = (f"Attendance Register  ·  {name}"
               + (f"  [{emp_id}]" if emp_id else "")
               + f"  ·  {loc}  ·  {emp['type']}  ·  {calendar.month_name[month]} {year}")
    t.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    t.fill = HDR_FILL; t.alignment = CENTER
    ws.row_dimensions[1].height = 24

    # Row 2: employee detail sub-header
    info_parts = [x for x in [desig, dept, area_v] if x]
    ws.merge_cells("A2:M2")
    info_cell = ws["A2"]
    info_cell.value = "  ·  ".join(info_parts) if info_parts else ""
    info_cell.font = Font(bold=False, color="FFFFFF", name="Calibri", size=9)
    info_cell.fill = _fill("2E4A7A"); info_cell.alignment = CENTER
    ws.row_dimensions[2].height = 16

    for ci, h in enumerate(["Date","Day","Status","Day Type","In Time","Out Time",
                             "Extra Punches","Extra Hours","Short Hours",
                             "Long Hours","Net Hours","Late Mark","Half Day"], 1):
        _hdr(ws.cell(row=3, column=ci), h)
    ws.row_dimensions[3].height = 18

    DR = 4
    if is_factory:
        sh_in  = '"09:15"'; sh_out = '"17:45"'
        lh_in  = '"08:30"'; lh_out = '"18:45"'
        lm_lo  = '"10:00"'; lm_hi  = '"13:30"'
        hd_thr = '"13:30"'
    else:
        sh_in  = '"09:45"'; sh_out = '"18:15"'
        lh_in  = '"09:15"'; lh_out = '"19:15"'
        lm_lo  = '"10:30"'; lm_hi  = '"14:00"'
        hd_thr = '"14:00"'

    for day_num in range(1, days_in_month + 1):
        r  = DR + day_num - 1
        d  = date(year, month, day_num)
        dt = _day_type(d, week_off, holidays)
        in_t, out_t = punch.get(day_num, (None, None))
        extra_punches = emp.get("extra_punches_by_day", {}).get(day_num, [])
        extra_mins = _extra_punch_mins(extra_punches)
        status = "Present" if (in_t or out_t) else "Absent"
        if extra_mins or any(p.get("in_time") or p.get("out_time") for p in extra_punches):
            status = "Present"

        dc = ws.cell(row=r, column=1)
        dc.value = d; dc.number_format = "DD-MMM-YY"
        dc.font = _font(); dc.alignment = CENTER; dc.border = THIN_B

        _dat(ws.cell(row=r, column=2), d.strftime("%a"))

        sc = ws.cell(row=r, column=3)
        sf = (WO_F if dt == "Week Off" else HOL_F if dt == "Company Holiday"
              else PRESENT_F if status == "Present" else ABSENT_F)
        _dat(sc, status, fill=sf)

        dtc = ws.cell(row=r, column=4)
        _dat(dtc, dt, fill=WO_F if dt=="Week Off" else (HOL_F if dt=="Company Holiday" else None))

        ec = ws.cell(row=r, column=5)
        fc = ws.cell(row=r, column=6)
        if in_t:  ec.value = _t(in_t);  ec.number_format = "HH:MM"
        if out_t: fc.value = _t(out_t); fc.number_format = "HH:MM"
        for cell in (ec, fc):
            cell.font = _font(); cell.alignment = CENTER; cell.border = THIN_B

        epc = ws.cell(row=r, column=7)
        epc.value = _extra_punch_label(extra_punches)
        epc.font = _font(); epc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        epc.border = THIN_B

        eh = ws.cell(row=r, column=8)
        eh.value = extra_mins / 1440.0 if extra_mins else 0
        eh.number_format = "[h]:mm"; eh.font = _font(); eh.alignment = CENTER; eh.border = THIN_B

        C = f"C{r}"; D = f"D{r}"; E = f"E{r}"; F = f"F{r}"

        sh = ws.cell(row=r, column=9)
        sh.value = (f'=IF(AND({D}="Working Day",{C}="Present",{E}<>"",{F}<>""),'
                    f'MAX(0,{E}-TIMEVALUE({sh_in}))+MAX(0,TIMEVALUE({sh_out})-{F}),0)')
        sh.number_format = "[h]:mm"; sh.font = _font(); sh.alignment = CENTER; sh.border = THIN_B

        lh = ws.cell(row=r, column=10)
        lh.value = (f'=IF(AND({D}="Working Day",{C}="Present",{E}<>"",{F}<>""),'
                    f'MAX(0,TIMEVALUE({lh_in})-{E})+MAX(0,{F}-TIMEVALUE({lh_out})),0)')
        lh.number_format = "[h]:mm"; lh.font = _font(); lh.alignment = CENTER; lh.border = THIN_B

        H = f"H{r}"; I = f"I{r}"; J = f"J{r}"
        nh = ws.cell(row=r, column=11)
        M = f"M{r}"
        nh.value = (f'=IF({M}="Half Day",0,'
                    f'IF(AND({D}="Working Day",{C}="Present",{E}<>"",{F}<>""),{J}-{I},'
                    f'IF(OR(AND({D}="Company Holiday",{C}="Present",{E}<>"",{F}<>""),'
                    f'AND({D}="Week Off",{C}="Present",{E}<>"",{F}<>"")),'
                    f'{F}-{E},0)))+{H}')
        nh.number_format = "[h]:mm"; nh.font = _font(); nh.alignment = CENTER; nh.border = THIN_B

        lm = ws.cell(row=r, column=12)
        lm.value = (f'=IF(AND({D}="Working Day",{C}="Present",{E}<>""),'
                    f'IF(AND({E}>TIMEVALUE({lm_lo}),{E}<TIMEVALUE({lm_hi})),"Late Mark",""),"")')
        lm.font = _font(); lm.alignment = CENTER; lm.border = THIN_B

        hd = ws.cell(row=r, column=13)
        hd.value = (f'=IF(AND({D}="Working Day",{C}="Present"),'
                    f'IF(OR({E}>=TIMEVALUE({hd_thr}),AND({F}<>"",{F}<=TIMEVALUE({hd_thr}))),"Half Day",""),"")')
        hd.font = _font(); hd.alignment = CENTER; hd.border = THIN_B
        ws.row_dimensions[r].height = 15

    # ── Monthly Summary ────────────────────────────────────────────────────────
    last = DR + days_in_month - 1
    sr   = last + 2
    CR = f"C{DR}:C{last}"; DR2 = f"D{DR}:D{last}"
    KR = f"M{DR}:M{last}"; IR  = f"K{DR}:K{last}"

    ws.merge_cells(f"A{sr}:B{sr}")
    _hdr(ws[f"A{sr}"], "MONTHLY SUMMARY")
    ws.row_dimensions[sr].height = 18

    ms_rows = {}
    for i, (lbl, frm) in enumerate([
        ("Total Working Days", f'=COUNTIF({DR2},"Working Day")'),
        ("Present Days",       f'=COUNTIFS({CR},"Present",{DR2},"Working Day")'),
        ("Absent Days",        f'=COUNTIFS({CR},"Absent",{DR2},"Working Day")'),
        ("Half Days",          f'=COUNTIF({KR},"Half Day")'),
        ("Total Net Hours",    f'=SUM({IR})'),
    ]):
        rr = sr + 1 + i
        ws.merge_cells(f"A{rr}:B{rr}")
        _lbl(ws.cell(row=rr, column=1), lbl)
        fc = ws.cell(row=rr, column=3)
        fc.value = frm
        if lbl == "Total Net Hours": fc.number_format = "[h]:mm"
        fc.font = _font(); fc.alignment = CENTER; fc.border = THIN_B
        ms_rows[lbl] = f"C{rr}"
        ws.row_dimensions[rr].height = 15

    # ── Leave Summary ──────────────────────────────────────────────────────────
    ls = sr + 7
    ws.merge_cells(f"A{ls}:B{ls}")
    _hdr(ws[f"A{ls}"], "LEAVE SUMMARY")
    ws.row_dimensions[ls].height = 18

    absent_ref  = ms_rows["Absent Days"]
    half_ref    = ms_rows["Half Days"]
    el_o = emp["el"] or 0; cl_o = emp["cl"] or 0
    sl_o = emp["sl"] or 0; sl_a = emp["sl_applied"] or 0

    r0 = ls + 1; r1 = ls + 2; r2 = ls + 3
    r4 = ls + 5; r5 = ls + 6
    r7 = ls + 8; r8 = ls + 9; r9 = ls + 10; r10 = ls + 11
    r11= ls + 12; r12= ls + 13; r13= ls + 14; r14= ls + 15
    r16= ls + 17; r17= ls + 18; r18= ls + 19

    leave_rows = [
        (r0,  "Opening EL Balance",            el_o,  None,      None),
        (r1,  "Opening CL Balance",            cl_o,  None,      None),
        (r2,  "Opening SL Balance",            sl_o,  None,      None),
        (r4,  "Cumulative Net Hours",           f'=SUM({IR})', "[h]:mm", None),
        (r5,  "EL Accrual Days",
              f'=IF(C{r4}<0,CEILING(C{r4}/TIMEVALUE("9:00:00"),1),'
              f'FLOOR(C{r4}/TIMEVALUE("9:00:00"),1))', None, None),
        (r7,  "Total Absent (Working Days)",    f"={absent_ref}+0.5*{half_ref}", None, None),
        (r8,  "CL Adjusted",                   f'=MIN(1,C{r1},MAX(0,C{r7}))', None, None),
        (r9,  "SL Adjusted",                   f'=MIN({sl_a},C{r2})', None, None),
        (r10, "Adjusted Opening EL Balance",   f'=MAX(0,C{r0}+C{r5})', None, None),
        (r11, "EL Adjusted (for Absences)",
              f'=MIN(C{r10},MAX(0,C{r7}-C{r8}-C{r9}))', None, None),
        (r12, "LWP from Negative EL Accrual",
              f'=IF(C{r0}+C{r5}<0,-1*(C{r0}+C{r5}),0)', None, None),
        (r13, "LWP from Excess Absences",
              f'=MAX(0,C{r7}-C{r8}-C{r9}-C{r11})', None, None),
        (r14, "Total LWP",                     f'=C{r12}+C{r13}', None, LWP_F),
        (r16, "Closing CL Balance",            f'=MAX(0,C{r1}-C{r8})', None, CLOSE_F),
        (r17, "Closing EL Balance",            f'=MAX(0,C{r10}-C{r11})', None, CLOSE_F),
        (r18, "Closing SL Balance",            f'=MAX(0,C{r2}-C{r9})', None, CLOSE_F),
    ]

    for rr, lbl, val, fmt, fill in leave_rows:
        ws.row_dimensions[rr].height = 15
        ws.merge_cells(f"A{rr}:B{rr}")
        _lbl(ws.cell(row=rr, column=1), lbl)
        vc = ws.cell(row=rr, column=3)
        vc.value = val
        if fmt:  vc.number_format = fmt
        if fill: vc.fill = fill
        vc.font = _font(bold=fill is not None)
        vc.alignment = CENTER; vc.border = THIN_B

    _set_widths(ws, {"A":28,"B":8,"C":12,"D":16,"E":10,"F":10,
                     "G":20,"H":12,"I":12,"J":12,"K":12,"L":12,"M":12})


# ── Labour sheet ──────────────────────────────────────────────────────────────
def _build_labour_sheet(wb, emp: dict, punch_data: dict, name_to_code: dict,
                        holidays: list, year: int, month: int):
    name     = emp["name"]
    week_off = emp["week_off"]
    ws       = wb.create_sheet(title=name[:28])
    ws.freeze_panes = "A4"

    code  = _resolve_code(name, name_to_code)
    punch = punch_data.get(code, {}) if code else {}
    days_in_month = calendar.monthrange(year, month)[1]

    emp_id  = emp.get("emp_id", "")
    desig   = emp.get("designation", "")
    dept    = emp.get("department", "")
    area_v  = emp.get("area", "")
    loc_l   = emp.get("location", "Factory")

    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = (f"Attendance Register  ·  {name}"
               + (f"  [{emp_id}]" if emp_id else "")
               + f"  ·  {loc_l}  ·  Labour  ·  {calendar.month_name[month]} {year}")
    t.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    t.fill = HDR_FILL; t.alignment = CENTER
    ws.row_dimensions[1].height = 24

    info_parts = [x for x in [desig, dept, area_v] if x]
    ws.merge_cells("A2:J2")
    info_cell = ws["A2"]
    info_cell.value = "  ·  ".join(info_parts) if info_parts else ""
    info_cell.font = Font(bold=False, color="FFFFFF", name="Calibri", size=9)
    info_cell.fill = _fill("2E4A7A"); info_cell.alignment = CENTER
    ws.row_dimensions[2].height = 16

    for ci, h in enumerate(["Date","Day","Status","Day Type",
                             "In Time","Out Time","Extra Punches","Extra Hours",
                             "Hours Worked","Overtime"], 1):
        _hdr(ws.cell(row=3, column=ci), h)
    ws.row_dimensions[3].height = 18

    DR = 4
    for day_num in range(1, days_in_month + 1):
        r  = DR + day_num - 1
        d  = date(year, month, day_num)
        dt = _day_type(d, week_off, holidays)
        in_t, out_t = punch.get(day_num, (None, None))
        extra_punches = emp.get("extra_punches_by_day", {}).get(day_num, [])
        extra_mins = _extra_punch_mins(extra_punches)
        status = "Present" if (in_t or out_t) else "Absent"
        if extra_mins or any(p.get("in_time") or p.get("out_time") for p in extra_punches):
            status = "Present"

        dc = ws.cell(row=r, column=1)
        dc.value = d; dc.number_format = "DD-MMM-YY"
        dc.font = _font(); dc.alignment = CENTER; dc.border = THIN_B

        _dat(ws.cell(row=r, column=2), d.strftime("%a"))

        sc = ws.cell(row=r, column=3)
        sf = (WO_F if dt == "Week Off" else HOL_F if dt == "Company Holiday"
              else PRESENT_F if status == "Present" else ABSENT_F)
        _dat(sc, status, fill=sf)

        dtc = ws.cell(row=r, column=4)
        _dat(dtc, dt, fill=WO_F if dt=="Week Off" else (HOL_F if dt=="Company Holiday" else None))

        ec = ws.cell(row=r, column=5)
        fc2 = ws.cell(row=r, column=6)
        if in_t:  ec.value = _t(in_t);  ec.number_format = "HH:MM"
        if out_t: fc2.value = _t(out_t); fc2.number_format = "HH:MM"
        for cell in (ec, fc2):
            cell.font = _font(); cell.alignment = CENTER; cell.border = THIN_B

        epc = ws.cell(row=r, column=7)
        epc.value = _extra_punch_label(extra_punches)
        epc.font = _font(); epc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        epc.border = THIN_B

        eh = ws.cell(row=r, column=8)
        eh.value = extra_mins / 1440.0 if extra_mins else 0
        eh.number_format = "[h]:mm"; eh.font = _font(); eh.alignment = CENTER; eh.border = THIN_B

        C = f"C{r}"; E = f"E{r}"; F = f"F{r}"

        H = f"H{r}"
        hw = ws.cell(row=r, column=9)
        hw.value = f'=IF(AND({C}="Present",{E}<>"",{F}<>""),{F}-{E},0)+{H}'
        hw.number_format = "[h]:mm"; hw.font = _font(); hw.alignment = CENTER; hw.border = THIN_B

        I = f"I{r}"
        ot = ws.cell(row=r, column=10)
        ot.value = f'=IF({C}="Present",{I}-TIMEVALUE("9:00:00"),0)'
        ot.number_format = "[h]:mm"; ot.font = _font(); ot.alignment = CENTER; ot.border = THIN_B
        ws.row_dimensions[r].height = 15

    last = DR + days_in_month - 1
    sr   = last + 2
    CR = f"C{DR}:C{last}"; DR2 = f"D{DR}:D{last}"; HR = f"J{DR}:J{last}"

    ws.merge_cells(f"A{sr}:B{sr}")
    _hdr(ws[f"A{sr}"], "MONTHLY SUMMARY")
    ws.row_dimensions[sr].height = 18

    for i, (lbl, frm) in enumerate([
        ("Working Days",        f'=COUNTIF({DR2},"Working Day")'),
        ("Present Days",        f'=COUNTIFS({CR},"Present",{DR2},"Working Day")'
                                f'+COUNTIFS({CR},"Present",{DR2},"Week Off")'
                                f'+COUNTIFS({CR},"Present",{DR2},"Company Holiday")'),
        ("Cumulative Overtime", f'=SUM({HR})'),
        ("OT Days",             f'=SUM({HR})/TIMEVALUE("6:00:00")'),
    ]):
        rr = sr + 1 + i
        ws.merge_cells(f"A{rr}:B{rr}")
        _lbl(ws.cell(row=rr, column=1), lbl)
        fc3 = ws.cell(row=rr, column=3)
        fc3.value = frm
        if lbl == "Cumulative Overtime": fc3.number_format = "[h]:mm"
        elif lbl == "OT Days":           fc3.number_format = "0.00"
        fc3.font = _font(); fc3.alignment = CENTER; fc3.border = THIN_B
        ws.row_dimensions[rr].height = 15

    _set_widths(ws, {"A":14,"B":8,"C":10,"D":16,"E":10,"F":10,
                     "G":20,"H":12,"I":12,"J":12})


# ── Rebuild from edited frontend JSON ─────────────────────────────────────────
def rebuild_from_json(emp_jsons: list, holidays: list,
                      output_path: str, year: int, month: int):
    """Rebuild Excel workbook using employee daily data from frontend (may include user edits)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for emp_json in emp_jsons:
        code = emp_json.get('code', '?')
        name = emp_json['name']
        # Reconstruct per-day punch dict from the daily array
        punch = {d['day']: (d.get('in_time'), d.get('out_time'))
                 for d in emp_json['daily']}
        extra_punches_by_day = {
            d['day']: d.get('extra_punches', [])
            for d in emp_json['daily']
        }
        punch_data_local  = {code: punch}
        name_to_code_local = {name.lower(): code}
        emp = {
            'name':        name,
            'emp_id':      emp_json.get('emp_id', ''),
            'designation': emp_json.get('designation', ''),
            'department':  emp_json.get('department', ''),
            'area':        emp_json.get('area', ''),
            'location':    emp_json['location'],
            'type':        emp_json['type'],
            'week_off':    emp_json.get('week_off', 6),
            'el':          emp_json.get('el_opening'),
            'cl':          emp_json.get('cl_opening'),
            'sl':          emp_json.get('sl_opening'),
            'sl_applied':  emp_json.get('sl_applied'),
            'extra_punches_by_day': extra_punches_by_day,
        }
        if emp['type'] == 'Labour':
            _build_labour_sheet(wb, emp, punch_data_local,
                                name_to_code_local, holidays, year, month)
        else:
            _build_employee_sheet(wb, emp, punch_data_local,
                                  name_to_code_local, holidays, year, month)

    wb.save(output_path)


def _parse_register_title(title: str) -> dict:
    parts = [p.strip() for p in str(title or "").split("·")]
    if len(parts) < 5 or "Attendance Register" not in parts[0]:
        raise ValueError("Workbook does not look like a generated attendance register")

    name_part = parts[1]
    emp_id = ""
    m = re.search(r"\[(.*?)\]", name_part)
    if m:
        emp_id = m.group(1).strip()
        name = re.sub(r"\s*\[.*?\]\s*", "", name_part).strip()
    else:
        name = name_part.strip()

    month_year = parts[4]
    month = year = None
    mm = re.match(r"([A-Za-z]+)\s+(\d{4})", month_year)
    if mm:
        month = list(calendar.month_name).index(mm.group(1)) if mm.group(1) in calendar.month_name else None
        year = int(mm.group(2))

    return {
        "name": name,
        "emp_id": emp_id,
        "location": parts[2],
        "type": parts[3],
        "month": month,
        "year": year,
    }

def _parse_register_info(info: str) -> dict:
    parts = [p.strip() for p in str(info or "").split("·")]
    return {
        "designation": parts[0] if len(parts) > 0 else "",
        "department": parts[1] if len(parts) > 1 else "",
        "area": parts[2] if len(parts) > 2 else "",
    }

def _find_label_value(ws, label: str):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if str(row[0].value or "").strip() == label:
            return row[2].value if len(row) >= 3 else None
    return None

def _parse_sl_applied(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"MIN\(([-+]?\d+(?:\.\d+)?)\s*,", str(v), re.IGNORECASE)
    return float(m.group(1)) if m else 0.0

def _week_off_from_daily(daily_rows: list, location: str) -> int:
    counts = {}
    for d in daily_rows:
        if d["day_type"] == "Week Off":
            counts[d["date_obj"].weekday()] = counts.get(d["date_obj"].weekday(), 0) + 1
    if counts:
        return max(counts, key=counts.get)
    return 1 if "factory" in location.lower() else 6

def parse_generated_register(xlsx_path: str) -> dict:
    """Parse a previously generated attendance register back into dashboard JSON."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    employees = []
    year = month = None
    holidays = set()

    for ws in wb.worksheets:
        headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        if headers[:6] != ["Date", "Day", "Status", "Day Type", "In Time", "Out Time"]:
            continue

        meta = _parse_register_title(ws["A1"].value)
        meta.update(_parse_register_info(ws["A2"].value))
        year = year or meta.get("year")
        month = month or meta.get("month")

        daily_rows = []
        for r in range(4, ws.max_row + 1):
            raw_date = ws.cell(r, 1).value
            if not isinstance(raw_date, (datetime, date)):
                break
            d_obj = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            day_type = ws.cell(r, 4).value or "Working Day"
            if day_type == "Company Holiday":
                holidays.add(d_obj)

            in_t = _cell_time_to_str(ws.cell(r, 5).value)
            out_t = _cell_time_to_str(ws.cell(r, 6).value)
            extra_punches = _parse_extra_punches(ws.cell(r, 7).value)
            metrics = _compute_day_row(meta["type"], in_t, out_t, day_type,
                                       meta["location"], extra_punches)
            status = ws.cell(r, 3).value or metrics["status"]
            metrics["status"] = status
            if ws.max_column >= 12:
                late_value = ws.cell(r, 12).value
                if isinstance(late_value, str) and not late_value.startswith("="):
                    metrics["late_mark"] = late_value
            if ws.max_column >= 13:
                half_value = ws.cell(r, 13).value
                if isinstance(half_value, str) and not half_value.startswith("="):
                    metrics["half_day"] = half_value

            daily_rows.append({
                "day": d_obj.day,
                "date": d_obj.strftime("%d-%b"),
                "date_obj": d_obj,
                "dow": ws.cell(r, 2).value or d_obj.strftime("%a"),
                "day_type": day_type,
                "in_time": in_t,
                "out_time": out_t,
                "extra_punches": extra_punches,
                **metrics,
            })

        if not daily_rows:
            continue

        week_off = _week_off_from_daily(daily_rows, meta["location"])
        emp = {
            "name": meta["name"],
            "code": meta["emp_id"] or "?",
            "emp_id": meta["emp_id"],
            "designation": meta["designation"],
            "department": meta["department"],
            "area": meta["area"],
            "location": meta["location"],
            "type": meta["type"],
            "week_off": week_off,
            "daily": [{k: v for k, v in row.items() if k != "date_obj"} for row in daily_rows],
        }

        if meta["type"] == "Labour":
            emp.update(_compute_emp_summary({"type": "Labour"}, emp["daily"]))
        else:
            emp["el_opening"] = float(_find_label_value(ws, "Opening EL Balance") or 0)
            emp["cl_opening"] = float(_find_label_value(ws, "Opening CL Balance") or 0)
            emp["sl_opening"] = float(_find_label_value(ws, "Opening SL Balance") or 0)
            emp["sl_applied"] = _parse_sl_applied(_find_label_value(ws, "SL Adjusted"))
            summary = _compute_emp_summary({
                "type": meta["type"],
                "el": emp["el_opening"],
                "cl": emp["cl_opening"],
                "sl": emp["sl_opening"],
                "sl_applied": emp["sl_applied"],
            }, emp["daily"])
            emp.update(summary)

        employees.append(emp)

    if not employees or not year or not month:
        raise ValueError("Could not parse generated attendance register workbook")

    return {
        "month": f"{calendar.month_name[month]} {year}",
        "year": year,
        "month_num": month,
        "total_emp": len(employees),
        "employees": employees,
        "holidays": [str(h) for h in sorted(holidays)],
    }


# ── Main entry ────────────────────────────────────────────────────────────────
def build_attendance_workbook(punch_data: dict, name_to_code: dict,
                              meta: dict, output_path: str,
                              year: int = None, month: int = None) -> dict:
    if year is None or month is None:
        from datetime import date as dt_today
        today = dt_today.today()
        year  = year  or today.year
        month = month or today.month

    holidays  = meta["holidays"]
    employees = meta["employees"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    emp_jsons = []
    for emp in employees:
        if emp["type"] == "Labour":
            _build_labour_sheet(wb, emp, punch_data, name_to_code, holidays, year, month)
        else:
            _build_employee_sheet(wb, emp, punch_data, name_to_code, holidays, year, month)
        emp_jsons.append(_build_emp_json(emp, punch_data, name_to_code, holidays, year, month))

    wb.save(output_path)

    return {
        "month":     f"{calendar.month_name[month]} {year}",
        "year":      year,
        "month_num": month,
        "total_emp": len(employees),
        "employees": emp_jsons,
        "holidays":  [str(h) for h in holidays],
    }
